from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn


SOURCE = r"D:\unity项目\LowPolyGame\李祺_Unity客户端开发简历_个人信息紧凑版.docx"
OUT = r"D:\unity项目\LowPolyGame\李祺_Unity客户端开发简历_课程技能优化版.docx"
BLUE = "24458D"
TEXT = "202735"
MUTED = "5B6473"


def set_font(run, size, color=TEXT, bold=False):
    run.font.name = "Microsoft YaHei"
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    rpr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    rpr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def add_field(p, label, value, size=9.5):
    r = p.add_run(label)
    set_font(r, size, BLUE, True)
    r = p.add_run(value)
    set_font(r, size)


def new_line(cell, after=0):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    return p


def clear_paragraphs(cell):
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)


doc = Document(SOURCE)

# Header: remove school address and use five evenly spaced field rows.
header_cell = doc.tables[0].cell(0, 0)
header_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
clear_paragraphs(header_cell)

p = new_line(header_cell, 5)
add_field(p, "姓名：", "李祺", 16.5)
p = new_line(header_cell, 7)
add_field(p, "性别：", "男", 10.0)
p.add_run("      ")
add_field(p, "出生年月：", "2005.01", 10.0)
p = new_line(header_cell, 7)
add_field(p, "意向岗位：", "Unity 客户端开发工程师", 10.0)
p = new_line(header_cell, 7)
add_field(p, "联系电话：", "15196885396", 10.0)
p = new_line(header_cell, 0)
add_field(p, "电子邮箱：", "2954046528@qq.com", 10.0)

# Education: retain the original school line and append major-course information.
education_cell = doc.tables[1].cell(0, 0)
p = education_cell.add_paragraph()
p.paragraph_format.space_before = Pt(3)
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 1.0
r = p.add_run("主要课程：")
set_font(r, 8.0, BLUE, True)
r = p.add_run("面向对象编程、算法与数据结构、软件工程；系统学习 C、C++。")
set_font(r, 8.0, MUTED)

# Skills: remove the nested two-column table and use one clean narrative column.
skills_cell = doc.tables[2].cell(0, 0)
for nested in list(skills_cell.tables):
    skills_cell._tc.remove(nested._tbl)
clear_paragraphs(skills_cell)

skill_items = [
    ("C#：", "熟悉面向对象、泛型、委托事件与协程机制，理解 GC 分配对运行时性能的影响。"),
    ("Unity 核心：", "熟悉 UGUI、Animator、Cinemachine、Input System、NavMesh 与 Tilemap，能够完成常见玩法和界面模块开发。"),
    ("工程化：", "了解 ScriptableObject 数据驱动、Addressables 异步加载与对象池复用，能够使用 Profiler 排查 GC 和 Draw Call 问题。"),
    ("架构设计：", "掌握 MVC、FSM、观察者模式、泛型单例与事件系统，能够基于 QFramework 进行模块分层和低耦合设计。"),
    ("工具协作：", "熟悉 Unity Editor 工具开发、Git/GitHub 协作及 NUnit EditMode 测试，具备基础的调试与代码验证能力。"),
]
for index, (title, description) in enumerate(skill_items):
    p = new_line(skills_cell, 2 if index < len(skill_items) - 1 else 0)
    add_field(p, title, description, 8.25)

# Self-assessment: add evidence from the internship without overstating responsibility.
self_cell = doc.tables[5].cell(0, 0)
clear_paragraphs(self_cell)
p = new_line(self_cell, 0)
summary = (
    "具备 Unity/C# 实际项目开发经验，参与过数字孪生和 2D、3D 游戏项目，熟悉 UI 交互、实体状态同步、数据驱动、"
    "游戏架构和 Unity 编辑器工具开发。工作认真负责，能够主动拆解问题并跟进功能闭环；在数字孪生实习中与后端协作完成"
    "接口数据接入、UI 展示与场景实体同步，具备良好的沟通协作意识和团队融入能力。"
)
r = p.add_run(summary)
set_font(r, 8.2)

doc.save(OUT)
print(OUT)
