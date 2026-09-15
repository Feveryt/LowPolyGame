from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE = r"D:\unity项目\LowPolyGame\李祺_Unity客户端开发简历_课程技能优化版.docx"
OUT = r"D:\unity项目\LowPolyGame\李祺_Unity客户端开发简历_顶部边距修复版.docx"


def set_margins(cell, top, start, bottom, end):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = margins.find(qn("w:" + side))
        if el is None:
            el = OxmlElement("w:" + side)
            margins.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


doc = Document(SOURCE)
header = doc.tables[0]
left = header.cell(0, 0)
photo = header.cell(0, 1)

# Explicit padding prevents text glyphs from touching the visible table border.
set_margins(left, top=150, start=145, bottom=135, end=95)
set_margins(photo, top=28, start=28, bottom=28, end=28)

# Ensure each field uses normal paragraph spacing instead of relying on an edge-adjacent line box.
for index, paragraph in enumerate(left.paragraphs):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(3 if index < len(left.paragraphs) - 1 else 0)
    paragraph.paragraph_format.line_spacing = 1.0

doc.save(OUT)
print(OUT)
