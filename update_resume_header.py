from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn


SOURCE = r"D:\unity项目\LowPolyGame\李祺_Unity客户端开发简历_技能描述版.docx"
OUT = r"D:\unity项目\LowPolyGame\李祺_Unity客户端开发简历_个人信息优化版.docx"
BLUE = "24458D"
TEXT = "202735"
MUTED = "5B6473"


def set_font(run, size, color, bold=False):
    run.font.name = "Microsoft YaHei"
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    rpr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    rpr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def add_field(paragraph, label, value, size=8.6):
    label_run = paragraph.add_run(label)
    set_font(label_run, size, BLUE, bold=True)
    value_run = paragraph.add_run(value)
    set_font(value_run, size, TEXT)


def make_line(cell, before=0, after=1):
    paragraph = cell.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0
    return paragraph


doc = Document(SOURCE)
header = doc.tables[0]
info_cell = header.cell(0, 0)
info_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

# Replace only the left information cell; the existing embedded photo remains untouched.
for paragraph in info_cell.paragraphs:
    paragraph._element.getparent().remove(paragraph._element)

p = make_line(info_cell, after=2)
add_field(p, "姓名：", "李祺", size=12.5)
p.add_run("      ")
add_field(p, "性别：", "男", size=9.0)
p.add_run("      ")
add_field(p, "出生年月：", "2005.01", size=9.0)

p = make_line(info_cell, after=1)
add_field(p, "意向岗位：", "Unity 客户端开发工程师", size=9.0)

p = make_line(info_cell, after=1)
add_field(p, "联系电话：", "15196885396", size=8.6)
p.add_run("      ")
add_field(p, "电子邮箱：", "2954046528@qq.com", size=8.6)

p = make_line(info_cell, after=0)
add_field(p, "学校地址：", "武汉工程大学流芳校区", size=8.6)

doc.save(OUT)
print(OUT)
