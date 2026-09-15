from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE


OUT = "李祯_Unity客户端开发简历.docx"
BLUE = "24458D"
LIGHT = "F6F8FB"
PALE_BLUE = "EAF0FF"
BORDER = "DDE4EF"
TEXT = "202735"
MUTED = "5B6473"


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=BORDER, size="6"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=110, bottom=80, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m)
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    tablePr = table._tbl.tblPr
    tblLayout = tablePr.find(qn("w:tblLayout"))
    if tblLayout is None:
        tblLayout = OxmlElement("w:tblLayout")
        tablePr.append(tblLayout)
    tblLayout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(width * 1440)))
            tcW.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_run_font(run, name="Microsoft YaHei", size=9, color=TEXT, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def style_paragraph(p, before=0, after=0, line=1.0, left=0, first=0):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.left_indent = Inches(left)
    fmt.first_line_indent = Inches(first)


def clear_cell(cell):
    if not cell.paragraphs:
        p = cell.add_paragraph()
    else:
        p = cell.paragraphs[0]
    p.clear()
    return p


def add_text(p, text, size=9, color=TEXT, bold=False, italic=False):
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return r


def add_bullet(cell, text, size=8.2, color=TEXT, before=0, after=1):
    p = cell.add_paragraph()
    style_paragraph(p, before=before, after=after, line=1.0, left=0.12, first=-0.12)
    add_text(p, "• ", size=size, color=BLUE, bold=True)
    add_text(p, text, size=size, color=color)
    return p


def add_section_heading(doc, title):
    p = doc.add_paragraph()
    style_paragraph(p, before=5, after=2, line=1.0)
    add_text(p, "◆ ", size=10.5, color=BLUE, bold=True)
    add_text(p, title, size=11.2, color=BLUE, bold=True)
    return p


def add_card(doc, padding=(85, 120, 85, 120)):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT)
    set_cell_border(cell)
    set_cell_margins(cell, *padding)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table, cell


def add_skill_row(table, label, content):
    row = table.add_row()
    row.cells[0].width = Inches(0.83)
    row.cells[1].width = Inches(5.95)
    label_cell, content_cell = row.cells
    set_cell_shading(label_cell, PALE_BLUE)
    set_cell_border(label_cell, PALE_BLUE, "0")
    set_cell_border(content_cell, LIGHT, "0")
    set_cell_margins(label_cell, 42, 75, 42, 75)
    set_cell_margins(content_cell, 42, 75, 42, 75)
    lp = clear_cell(label_cell)
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(lp, line=1.0)
    add_text(lp, label, size=8.2, color=BLUE, bold=True)
    cp = clear_cell(content_cell)
    style_paragraph(cp, line=1.0)
    add_text(cp, content, size=8.2, color=TEXT)


def add_project_line(cell, name, meta, bullets):
    p = cell.add_paragraph()
    style_paragraph(p, before=1, after=0, line=1.0)
    add_text(p, name, size=8.7, color=TEXT, bold=True)
    add_text(p, "  |  " + meta, size=8.0, color=MUTED)
    for b in bullets:
        add_bullet(cell, b, size=8.0, after=0)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.34)
section.bottom_margin = Inches(0.32)
section.left_margin = Inches(0.48)
section.right_margin = Inches(0.48)
section.header_distance = Inches(0.15)
section.footer_distance = Inches(0.15)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Microsoft YaHei"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(8.6)
normal.font.color.rgb = RGBColor.from_string(TEXT)

# Header information
top = doc.add_table(rows=1, cols=2)
top.alignment = WD_TABLE_ALIGNMENT.CENTER
top.autofit = False
set_table_widths(top, [6.05, 1.10])
left, photo = top.rows[0].cells
set_cell_margins(left, 0, 0, 0, 0)
set_cell_margins(photo, 0, 0, 0, 0)
left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
photo.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

p = clear_cell(left)
style_paragraph(p, after=1, line=1.0)
add_text(p, "李祯", size=23, color=BLUE, bold=True)
p = left.add_paragraph()
style_paragraph(p, after=1, line=1.0)
add_text(p, "Unity 客户端开发工程师", size=9.3, color=BLUE, bold=True)
p = left.add_paragraph()
style_paragraph(p, after=1, line=1.0)
add_text(p, "男    |    2005.01    |    15196885396", size=8.5, color=MUTED)
p = left.add_paragraph()
style_paragraph(p, after=1, line=1.0)
add_text(p, "2954046528@qq.com    |    武汉工程大学流芳校区", size=8.5, color=MUTED)

set_cell_shading(photo, "E9EEF8")
set_cell_border(photo, "B8C7E2", "8")
set_cell_margins(photo, 0, 0, 0, 0)
pp = clear_cell(photo)
pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
style_paragraph(pp, line=1.0)
add_text(pp, "寸照\n点击替换", size=7.5, color=BLUE, bold=True)
photo.width = Inches(1.10)
top.rows[0].height = Inches(1.38)
top.rows[0].height_rule = 1

add_section_heading(doc, "教育背景")
edu, cell = add_card(doc, padding=(65, 120, 65, 120))
set_table_widths(edu, [7.15])
p = clear_cell(cell)
style_paragraph(p, line=1.0)
add_text(p, "2023.09 - 2027.06", size=8.5, color=MUTED)
add_text(p, "                                      武汉工程大学", size=9.0, color=TEXT, bold=True)
add_text(p, "                                      数字媒体技术｜本科", size=8.5, color=MUTED)

add_section_heading(doc, "专业技能")
skills, cell = add_card(doc, padding=(70, 110, 70, 110))
set_table_widths(skills, [7.15])
# 删除卡片占位段落，避免技能表格上方出现不必要留白。
cell._tc.remove(cell.paragraphs[0]._p)
inner = cell.add_table(rows=0, cols=2)
inner.alignment = WD_TABLE_ALIGNMENT.CENTER
inner.autofit = False
set_table_widths(inner, [1.22, 5.82])
add_skill_row(inner, "C#", "面向对象、委托事件、泛型、协程、GC")
add_skill_row(inner, "Unity 核心", "UGUI、Animator、Cinemachine、Input System、NavMesh、Tilemap")
add_skill_row(inner, "工程化", "ScriptableObject、Addressables、对象池、Profiler、Draw Call 优化")
add_skill_row(inner, "架构设计", "MVC、FSM、观察者模式、泛型单例、事件系统、QFramework 分层")
add_skill_row(inner, "工具协作", "Unity Editor、Git、GitHub、NUnit")

add_section_heading(doc, "实习经历")
intern, cell = add_card(doc, padding=(75, 120, 75, 120))
set_table_widths(intern, [7.15])
p = clear_cell(cell)
style_paragraph(p, line=1.0, after=0)
add_text(p, "XX科技有限公司", size=8.8, color=TEXT, bold=True)
add_text(p, "  |  Unity 开发实习生  |  20XX.XX - 20XX.XX", size=8.3, color=MUTED)
add_bullet(cell, "参与 Unity + C# 数字孪生项目，负责后端接口接入、数据解析及 UI/三维场景展示。", size=8.0)
add_bullet(cell, "根据接口下发的楼层数据同步电梯模型状态，通过插值实现模型向目标楼层平滑移动。", size=8.0)
add_bullet(cell, "制定业务坐标到 Unity 世界坐标的映射规则，完成车辆位置、轨迹生成及车辆沿轨迹运动。", size=8.0)
add_bullet(cell, "负责考试管理系统 UGUI 搭建与面板切换，实现考试选择、学生/班组筛选及新增人员逻辑。", size=8.0)

add_section_heading(doc, "项目经历")
projects, cell = add_card(doc, padding=(70, 120, 70, 120))
set_table_widths(projects, [7.15])
cell._tc.remove(cell.paragraphs[0]._p)
add_project_line(cell, "Low Poly ARPG", "Unity 客户端开发｜个人项目", [
    "基于 QFramework 按 Controller / Command / System / Model 分层，完成第三人称移动战斗、任务、背包、商店和 HUD 等模块。",
    "以 DialogueAsset 驱动 NPC 台词、分支、任务条件与事件；开发对话编辑器，支持资产创建、节点/选项编辑、Undo、引用清理及循环/不可达节点校验。",
    "开发 Camera Sequence Recorder，从 Scene View 采样关键帧；运行时使用 Cinemachine 播放静态演出和玩家-NPC 双人构图，并处理输入锁定与镜头回退。",
    "使用 ScriptableObject 进行配置，编写 EditMode 测试覆盖任务、伤害、资源和音效等核心逻辑。",
])
add_project_line(cell, "《勇士传说》", "2D 横版动作冒险｜个人项目", [
    "使用 FSM 实现敌人巡逻、追击、攻击等行为；结合 Addressables、Input System 与 ScriptableObject 事件完成场景和玩法模块解耦。",
])
add_project_line(cell, "3D RPG 原型", "Unity 角色扮演｜个人项目", [
    "基于 NavMeshAgent 实现点击移动、自动接近攻击和敌人 AI；使用 ScriptableObject 配置属性，完成战斗、升级、传送与存档。",
])
add_project_line(cell, "弹幕射击 STG", "Unity 2D｜个人项目", [
    "基于 Unity ObjectPool<T> 复用子弹与特效，使用 XML 配置表驱动 5 类弹道和发射波次，并封装泛型 UI 面板基类。",
])

add_section_heading(doc, "自我评价")
self_card, cell = add_card(doc, padding=(75, 120, 75, 120))
set_table_widths(self_card, [7.15])
p = clear_cell(cell)
style_paragraph(p, line=1.05)
add_text(p, "具备 Unity/C# 实际项目开发经验，参与过数字孪生和 2D、3D 游戏项目，熟悉 UI 交互、实体状态同步、数据驱动、游戏架构和 Unity 编辑器工具开发。能够从需求拆解、功能实现到问题定位完成模块闭环，并持续通过测试和重构提升代码稳定性。", size=8.2, color=TEXT)

# Footer with a restrained editable note
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
style_paragraph(fp, line=1.0)
add_text(fp, "李祯｜Unity 客户端开发", size=7.2, color="8A94A6")

doc.save(OUT)
print(OUT)
