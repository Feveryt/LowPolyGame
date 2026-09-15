from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn


PATH = r"D:\unity项目\LowPolyGame\李祺_Unity客户端开发简历.docx"
OUT = r"D:\unity项目\LowPolyGame\李祺_Unity客户端开发简历_技能描述版.docx"
TEXT = "202735"

SKILL_ROWS = [
    ("C#", "熟悉 C# 面向对象、泛型、委托事件与协程机制，理解 GC 分配对运行时性能的影响。"),
    ("Unity 核心", "熟悉 UGUI、Animator、Cinemachine、Input System、NavMesh 与 Tilemap，能够完成常见玩法和界面模块开发。"),
    ("工程化", "了解 ScriptableObject 数据驱动、Addressables 异步加载与对象池复用，能够使用 Profiler 排查 GC 和 Draw Call 问题。"),
    ("架构设计", "掌握 MVC、FSM、观察者模式、泛型单例与事件系统，能够基于 QFramework 进行模块分层和低耦合设计。"),
    ("工具协作", "熟悉 Unity Editor 工具开发、Git/GitHub 协作及 NUnit EditMode 测试，具备基础的调试与代码验证能力。"),
]


def set_font(run, size=8.2, color=TEXT, bold=False):
    run.font.name = "Microsoft YaHei"
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    rpr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    rpr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def replace_cell_text(cell, text, size=8.2, color=TEXT, bold=False, center=False):
    paragraph = cell.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    if center:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_font(run, size=size, color=color, bold=bold)


doc = Document(PATH)
skills_table = doc.tables[2].cell(0, 0).tables[0]
if len(skills_table.rows) != len(SKILL_ROWS):
    raise RuntimeError(f"Expected {len(SKILL_ROWS)} skill rows, found {len(skills_table.rows)}")

for row, (label, description) in zip(skills_table.rows, SKILL_ROWS):
    replace_cell_text(row.cells[0], label, size=8.2, color="24458D", bold=True, center=True)
    replace_cell_text(row.cells[1], description, size=8.15)

doc.save(OUT)
print(OUT)
