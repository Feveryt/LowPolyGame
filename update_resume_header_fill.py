from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn


SOURCE = r"D:\unity项目\LowPolyGame\李祺_Unity客户端开发简历_个人信息优化版.docx"
OUT = r"D:\unity项目\LowPolyGame\李祺_Unity客户端开发简历_个人信息紧凑版.docx"
BLUE = "24458D"
TEXT = "202735"


def font(run, size, color=TEXT, bold=False):
    run.font.name = "Microsoft YaHei"
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), "Microsoft YaHei")
    rpr.rFonts.set(qn("w:hAnsi"), "Microsoft YaHei")
    rpr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def field(p, label, value, size=9.4):
    r = p.add_run(label)
    font(r, size, BLUE, True)
    r = p.add_run(value)
    font(r, size)


def line(cell, after):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    return p


doc = Document(SOURCE)
header = doc.tables[0]
cell = header.cell(0, 0)
cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

# Six lines deliberately occupy the same vertical rhythm as the 1-inch photo.
for p in cell.paragraphs:
    p._element.getparent().remove(p._element)

p = line(cell, 3)
field(p, "姓名：", "李祺", 14.5)
p = line(cell, 3)
field(p, "性别：", "男")
p.add_run("      ")
field(p, "出生年月：", "2005.01")
p = line(cell, 3)
field(p, "意向岗位：", "Unity 客户端开发工程师")
p = line(cell, 3)
field(p, "联系电话：", "15196885396")
p = line(cell, 3)
field(p, "电子邮箱：", "2954046528@qq.com")
p = line(cell, 0)
field(p, "学校地址：", "武汉工程大学流芳校区")

doc.save(OUT)
print(OUT)
